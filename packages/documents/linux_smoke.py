#!/usr/bin/env python3
"""Real Linux document smoke: python packages/documents/linux_smoke.py.

Requires the installed hr-documents package, Tesseract language packs, LibreOffice
and fonts-noto-cjk. No pytest/reportlab/model/API/real CV is used. All artifacts
live in a TemporaryDirectory removed on success or failure. Failures exit 1;
non-Linux execution fails explicitly instead of silently passing/skipping.
"""
import hashlib
import json
import math
from pathlib import Path
import sys
import tempfile


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def normalized(text):
    # Test matching only: allow common script variants, never modify parser output.
    return ''.join(text.split()).casefold().translate(str.maketrans('數據資', '数据资'))


def validate(result, label, method):
    require(result['status'] == 'ready', f'{label}: {result["status"]}: {result["warnings"]}')
    require(not result['warnings'], f'{label}: unexpected warnings: {result["warnings"]}')
    require(len(result['pages']) == 1 and result['blocks'], f'{label}: missing page/evidence')
    page = result['pages'][0]
    require(page['number'] == 1 and page['width'] > 0 and page['height'] > 0,
            f'{label}: invalid page dimensions')
    from PIL import Image
    with Image.open(page['imagePath']) as preview:
        preview.verify()
    for block in result['blocks']:
        require(block['page'] == 1 and block['method'] == method, f'{label}: wrong evidence method/page')
        box = block['box']
        require(box and all(math.isfinite(v) and 0 <= v <= 1 for v in box.values()),
                f'{label}: missing/invalid normalized box')
        require(box['width'] > 0 and box['height'] > 0 and
                box['x'] + box['width'] <= 1.000001 and
                box['y'] + box['height'] <= 1.000001, f'{label}: box outside page')
        if method == 'ocr':
            confidence = block['confidence']
            require(isinstance(confidence, (float, int)) and math.isfinite(confidence)
                    and 0 <= confidence <= 1, f'{label}: missing/invalid real OCR confidence')


def run():
    require(sys.platform.startswith('linux'), 'This is a real Linux smoke check; run it inside the deployment image.')
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from hr_documents import parse_document  # Public contract; there is no parse_file API.

    fonts = list(Path('/usr/share/fonts').rglob('NotoSansCJK-Regular.ttc'))
    require(bool(fonts), 'NotoSansCJK-Regular.ttc missing: install fonts-noto-cjk.')
    font = ImageFont.truetype(str(fonts[0]), 60)
    reports = []
    with tempfile.TemporaryDirectory(prefix='hr-documents-linux-smoke-') as temporary:
        folder = Path(temporary)
        image_path = folder / 'fictional multilingual.png'
        image = Image.new('RGB', (1800, 900), 'white')
        draw = ImageDraw.Draw(image)
        lines = [('FICTIONAL CV Python evidence', 'Python evidence', 100),
                 ('虚构简历 数据分析', '数据分析', 350),
                 ('虛構履歷 資料分析', '資料分析', 600)]
        expected = []
        for text, phrase, y in lines:
            draw.text((80, y), text, font=font, fill='black', anchor='lt')
            expected.append((phrase, draw.textbbox((80, y), text, font=font, anchor='lt')))
        image.save(image_path)
        before = hashlib.sha256(image_path.read_bytes()).hexdigest()
        result = parse_document(image_path, folder / 'previews', 'linux-image')
        validate(result, 'multilingual image', 'ocr')
        for phrase, (_, top, _, bottom) in expected:
            candidates = [b for b in result['blocks'] if normalized(phrase) in normalized(b['text'])]
            require(candidates, f'OCR missing {phrase!r}; actual text: {result["text"]!r}')
            center = (top + bottom) / 2 / image.height
            require(any(abs(b['box']['y'] + b['box']['height']/2 - center) < .035
                        for b in candidates), f'OCR {phrase!r} points to the wrong image line')
        require(hashlib.sha256(image_path.read_bytes()).hexdigest() == before, 'OCR modified source image')
        reports.append({'case': 'English + simplified/traditional Chinese image', 'status': 'pass',
                        'blocks': len(result['blocks']), 'ocrText': result['text']})

        pdf_path = folder / 'fictional-native.pdf'
        with fitz.open() as pdf:
            page = pdf.new_page(width=600, height=800)
            page.insert_text((60, 80), 'FICTIONAL CV - native PDF smoke', fontsize=18)
            page.insert_text((60, 180), 'Python evidence from fictional project', fontsize=16)
            pdf.save(pdf_path)
        result = parse_document(pdf_path, folder / 'previews', 'linux-native')
        validate(result, 'native PDF', 'native')
        block = next((b for b in result['blocks'] if 'Python evidence' in b['text']), None)
        require(block is not None, 'Native PDF source text missing')
        box = block['box']
        with fitz.open(pdf_path) as pdf:
            rect = fitz.Rect(box['x']*600, box['y']*800,
                             (box['x']+box['width'])*600, (box['y']+box['height'])*800)
            rect += (-.5, -.5, .5, .5)
            require('Python evidence' in pdf[0].get_text(clip=rect), 'Native PDF box does not locate source text')
        reports.append({'case': 'native PDF and exact source rectangle', 'status': 'pass'})

        docx_path = folder / 'fictional.docx'
        document = Document()
        document.add_paragraph('FICTIONAL CV - Linux DOCX smoke')
        document.add_paragraph('Python document workflow evidence')
        document.save(docx_path)
        before = hashlib.sha256(docx_path.read_bytes()).hexdigest()
        result = parse_document(docx_path, folder / 'previews', 'linux-docx')
        validate(result, 'DOCX', 'native')
        require(result['previewKind'] == 'rendered', 'DOCX fell back to text; real LibreOffice rendering required')
        require('Python document workflow evidence' in result['text'], 'Rendered DOCX source text missing')
        require(hashlib.sha256(docx_path.read_bytes()).hexdigest() == before, 'Renderer modified original DOCX')
        reports.append({'case': 'DOCX actual LibreOffice render and preview', 'status': 'pass'})
    # Printed only after TemporaryDirectory cleanup succeeded.
    print(json.dumps({'status': 'pass', 'platform': sys.platform, 'cases': reports,
                      'temporaryFilesRemoved': True, 'modelCalls': 0}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    try:
        run()
    except Exception as exc:
        print(f'FAIL Linux document smoke: {type(exc).__name__}: {exc}', file=sys.stderr)
        sys.exit(1)
