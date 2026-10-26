"""Synthetic regression for successful conversions that clip source paragraphs."""
from pathlib import Path
import pytest


@pytest.mark.parametrize('complete_render',[False,True])
def test_docx_successful_but_clipped_render_keeps_complete_source(tmp_path,monkeypatch,complete_render):
    from docx import Document
    from hr_documents import parser
    source=tmp_path/'source.docx';document=Document()
    first='Built a demonstration system with a team.'
    tail='Languages: English (Fluent), Mandarin (Native) and Cantonese (Native)'
    document.add_paragraph(first);document.add_paragraph(tail);document.save(source)
    renderer=tmp_path/'fake-renderer';renderer.write_text('not executed')
    monkeypatch.setattr(parser,'_renderer_path',lambda:renderer)
    def convert(command,**kwargs):
        out=Path(command[command.index('--outdir')+1]);(out/(source.stem+'.pdf')).write_bytes(b'fake')
    monkeypatch.setattr(parser.subprocess,'run',convert)
    def rendered(path,out,result):
        result['pages'].append({'number':1,'width':100,'height':100,'imagePath':str(out/'page.png')})
        parser._add(result,1,first,{'x':0,'y':0,'width':1,'height':.1})
        if complete_render:
            parser._add(result,1,tail.replace('Fluent','Fl uent'),{'x':0,'y':.2,'width':1,'height':.1})
    monkeypatch.setattr(parser,'_pdf',rendered)
    out=tmp_path/'out';out.mkdir();result={'id':'synthetic','pages':[],'blocks':[],'warnings':[],'status':'ready'}
    parser._docx(source,out,result)
    assert result['status']=='ready'
    if complete_render:
        assert result['previewKind']=='rendered' and result['pages'][0]['imagePath']
        assert not result['warnings']
    else:
        assert result['previewKind']=='text' and result['pages'][0]['imagePath'] is None
        assert [b['text'] for b in result['blocks']]==[first,tail]
        assert all(b['box'] is None for b in result['blocks'])
        assert any('预览不完整' in warning for warning in result['warnings'])
